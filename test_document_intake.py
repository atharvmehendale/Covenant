"""
Automatic checks for document intake — the PDF / Word / text reading layer.

Why this matters: the dangerous failure here is not a crash, it's SILENCE. If a
scanned contract extracts to an empty string and we hand that to the reviewer
agents, they will produce a confident, well-written review of nothing at all,
and a human might sign on the strength of it. These tests pin down that the
unreadable cases are reported as unreadable.

They also cover the awkward middle case a real firm hits: a document where some
pages have text and some are scanned images. That one is allowed through, but it
must carry a warning naming the pages nobody read.

No AI model is called and no network is used, so this runs offline in a second.
It builds its own PDFs on the fly in a temporary folder, so it is testing real
PDF parsing rather than a stand-in.

Run it with: python test_document_intake.py
It prints "ALL TESTS PASSED" or names the exact check that failed.
"""

import os
import shutil
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from document_intake import (
    MIN_TOTAL_CHARS,
    SCANNED_MESSAGE,
    UnreadableDocument,
    describe,
    extract,
    load_text,
    load_text_or_empty,
    normalize_text,
)

ROOT = os.path.dirname(os.path.abspath(__file__))
SAMPLES = os.path.join(ROOT, "samples")

# The real clause from the project's own sample contract. Every PDF test below
# checks that this exact wording survives the round trip through a PDF, because
# a review is worthless if the text it read wasn't the text on the page.
CLAUSE = (
    "'Operational Net-Zero' shall be defined exclusively as the mitigation of "
    "Scope 1 and Scope 2 emissions"
)

WORKSPACE = None  # a temp folder, created in main()


def one_line(text: str) -> str:
    """Collapses whitespace so a quote can be compared across line wraps."""
    return " ".join((text or "").split())


def make_pdf(name: str, paragraphs: list, scanned: bool = False) -> str:
    """
    Builds a genuine PDF in the temp workspace.

    `paragraphs` is a list of page contents — one entry per page. An entry of
    None makes a page that is an IMAGE only, which is how the partly-scanned
    case is built.
    """
    import fitz

    document = fitz.open()
    for body in paragraphs:
        page = document.new_page(width=595, height=842)
        if body is None:
            # A blank image page: no text layer, exactly like a scan.
            blank = fitz.open()
            source = blank.new_page(width=595, height=842)
            source.insert_text((72, 100), "scanned image content", fontsize=11)
            pixmap = source.get_pixmap(dpi=72)
            page.insert_image(page.rect, stream=pixmap.tobytes("jpeg", jpg_quality=50))
            blank.close()
        else:
            page.insert_textbox(fitz.Rect(72, 72, 523, 770), body, fontsize=11, fontname="helv")

    path = os.path.join(WORKSPACE, name)

    if scanned:
        # Re-render every page as an image, dropping the text layer entirely.
        flat = fitz.open()
        for page in document:
            pixmap = page.get_pixmap(dpi=110)
            target = flat.new_page(width=page.rect.width, height=page.rect.height)
            target.insert_image(target.rect, stream=pixmap.tobytes("jpeg", jpg_quality=55))
        flat.save(path)
        flat.close()
    else:
        document.save(path)
    document.close()
    return path


# ----------------------------------------------------------------------------
# The happy paths
# ----------------------------------------------------------------------------

def test_single_page_pdf_is_read():
    path = make_pdf("one_page.pdf", [f"Section 14.1 (Environmental Indemnity): {CLAUSE}."])
    result = extract(path)
    assert result["ok"] is True, f"A normal PDF was rejected: {result['message']}"
    assert result["page_count"] == 1, f"Expected 1 page, got {result['page_count']}"
    assert result["method"] == "pypdf", f"Expected pypdf, got {result['method']}"
    assert CLAUSE in one_line(result["text"]), "The clause did not survive extraction"


def test_multi_page_pdf_reads_every_page():
    """The one that would silently break: text only on the LAST page."""
    path = make_pdf(
        "three_pages.pdf",
        [
            "ARTICLE I - DEFINITIONS. " + ("Filler wording for page one. " * 30),
            "ARTICLE II - GOVERNANCE. " + ("Filler wording for page two. " * 30),
            f"ARTICLE XIV - ENVIRONMENTAL MATTERS. Section 14.1: {CLAUSE}.",
        ],
    )
    result = extract(path)
    assert result["ok"] is True, f"A multi-page PDF was rejected: {result['message']}"
    assert result["page_count"] == 3, f"Expected 3 pages, got {result['page_count']}"
    assert result["pages_with_text"] == 3, (
        f"Only {result['pages_with_text']} of 3 pages produced text"
    )
    text = one_line(result["text"])
    assert "ARTICLE I" in text, "Page 1 was not read"
    assert "page two" in text, "Page 2 was not read"
    assert CLAUSE in text, "Page 3 was not read — a clause deep in the document would be missed"


def test_plain_text_still_works():
    """The formats that worked before this feature must keep working unchanged."""
    path = os.path.join(WORKSPACE, "clause.txt")
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"Section 14.1 (Environmental Indemnity): {CLAUSE}.")
    result = extract(path)
    assert result["ok"] is True, f"A plain text file was rejected: {result['message']}"
    assert result["format"] == "plain text", f"Wrong format: {result['format']}"
    assert CLAUSE in one_line(result["text"]), "The clause did not survive a text read"


def test_non_utf8_text_file_is_still_read():
    """A .txt saved by a Windows program is often cp1252, not UTF-8."""
    path = os.path.join(WORKSPACE, "cp1252.txt")
    with open(path, "wb") as f:
        f.write(("Section 3.1 (Fees): the fee is £250,000 — payable monthly. " * 4).encode("cp1252"))
    result = extract(path)
    assert result["ok"] is True, f"A cp1252 text file was rejected: {result['message']}"
    assert "250,000" in result["text"], "The cp1252 file was read but its content is wrong"


def test_docx_including_table_text():
    """Contract terms very often live in a table, so tables must be read too."""
    try:
        import docx
    except ImportError:
        return  # python-docx isn't installed; the format is optional
    path = os.path.join(WORKSPACE, "terms.docx")
    document = docx.Document()
    document.add_paragraph(f"Section 22.3 (Reporting Scope): {CLAUSE}.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Scope 3 downstream"
    table.rows[0].cells[1].text = "Not reported under this Agreement"
    document.save(path)

    result = extract(path)
    assert result["ok"] is True, f"A Word document was rejected: {result['message']}"
    assert CLAUSE in one_line(result["text"]), "The clause did not survive a .docx read"
    assert "Not reported under this Agreement" in one_line(result["text"]), (
        "Text inside a table was skipped — table clauses would be missed"
    )


# ----------------------------------------------------------------------------
# The failures that must be reported, not swallowed
# ----------------------------------------------------------------------------

def test_scanned_pdf_is_detected_and_not_reviewed():
    """THE important one: an image-only PDF must never come back as readable."""
    path = make_pdf(
        "scanned.pdf",
        [f"Section 14.1 (Environmental Indemnity): {CLAUSE}."],
        scanned=True,
    )
    result = extract(path)
    assert result["ok"] is False, (
        "A scanned PDF was reported as readable — the agents would have "
        "reviewed an empty document"
    )
    assert result["problem"] == "scanned", f"Wrong problem code: {result['problem']}"
    assert SCANNED_MESSAGE in result["message"], (
        f"The message doesn't explain the scanned case: {result['message']}"
    )
    assert "OCR" in result["message"], "The message should say OCR isn't built yet"
    assert result["page_count"] == 1, "A scanned PDF should still report its page count"
    assert result["pages_with_text"] == 0, "A scanned PDF should report no pages with text"


def test_scanned_pdf_raises_rather_than_returning_empty_text():
    """load_text() is what the agents call, so it must refuse, not return ""."""
    path = make_pdf("scanned2.pdf", ["Some real wording on the page. " * 20], scanned=True)
    try:
        load_text(path)
    except UnreadableDocument as e:
        assert e.problem == "scanned", f"Wrong problem code on the exception: {e.problem}"
        return
    raise AssertionError("load_text() returned text for a scanned PDF instead of refusing")


def test_partly_scanned_pdf_is_reviewed_but_warns():
    """
    The awkward real-world middle: page 1 is a scan, the rest is text.

    This is allowed through — refusing it would be worse — but the result must
    say which pages nobody read.
    """
    path = make_pdf(
        "part_scanned.pdf",
        [None, f"ARTICLE XIV. Section 14.1: {CLAUSE}. " + ("Additional terms follow. " * 20)],
    )
    result = extract(path)
    assert result["ok"] is True, f"A partly-scanned PDF was refused outright: {result['message']}"
    assert result["page_count"] == 2, f"Expected 2 pages, got {result['page_count']}"
    assert result["pages_with_text"] == 1, (
        f"Expected 1 of 2 pages to have text, got {result['pages_with_text']}"
    )
    assert result["warnings"], "A partly-scanned document produced no warning at all"
    warning = " ".join(result["warnings"])
    assert "1" in warning and "2" in warning, f"The warning doesn't name the pages: {warning}"
    assert CLAUSE in one_line(result["text"]), "The readable page's clause was lost"


def test_tiny_amount_of_text_is_treated_as_unreadable():
    """
    A page bearing only a scanner stamp is not a contract.

    Extraction that returns a handful of characters must not be passed off as a
    successful read, or the reviewer agents get near-nothing and invent the rest.
    """
    path = make_pdf("stamp_only.pdf", ["RECEIVED 12 MAR"])
    result = extract(path)
    assert result["ok"] is False, (
        f"{MIN_TOTAL_CHARS}-character floor not enforced: a near-empty PDF was "
        "reported as readable"
    )
    assert result["problem"] == "scanned", f"Wrong problem code: {result['problem']}"
    assert "too little" in result["message"].lower(), (
        f"The message should say there wasn't enough text: {result['message']}"
    )


def test_damaged_pdf_does_not_crash():
    """A truncated or fake PDF must produce a message, never a traceback."""
    path = os.path.join(WORKSPACE, "broken.pdf")
    with open(path, "wb") as f:
        f.write(b"%PDF-1.7\nthis is not really a pdf at all\n%%EOF\n")
    result = extract(path)  # must not raise
    assert result["ok"] is False, "A corrupt PDF was reported as readable"
    assert result["problem"] in ("damaged", "scanned"), f"Wrong problem code: {result['problem']}"
    assert result["message"], "A corrupt PDF produced no explanation"


def test_password_protected_pdf_says_so():
    """Being locked is a different problem from being blank, and must read that way."""
    import fitz

    plain = make_pdf("to_lock.pdf", [f"Section 14.1: {CLAUSE}. " + ("Terms. " * 30)])
    locked = os.path.join(WORKSPACE, "locked.pdf")
    document = fitz.open(plain)
    document.save(
        locked,
        encryption=fitz.PDF_ENCRYPT_AES_256,
        owner_pw="owner-secret",
        user_pw="user-secret",
    )
    document.close()

    result = extract(locked)
    assert result["ok"] is False, "A password-protected PDF was reported as readable"
    assert result["problem"] == "encrypted", f"Wrong problem code: {result['problem']}"
    assert "password" in result["message"].lower(), (
        f"The message doesn't mention the password: {result['message']}"
    )


def test_unsupported_format_is_refused_by_name():
    path = os.path.join(WORKSPACE, "contract.rtf")
    with open(path, "w", encoding="utf-8") as f:
        f.write("{\\rtf1 some rtf content}")
    result = extract(path)
    assert result["ok"] is False, "An .rtf file was accepted"
    assert result["problem"] == "unsupported", f"Wrong problem code: {result['problem']}"
    assert ".rtf" in result["message"], f"The message should name the format: {result['message']}"


def test_empty_file_is_refused():
    path = os.path.join(WORKSPACE, "empty.pdf")
    open(path, "wb").close()
    result = extract(path)
    assert result["ok"] is False, "A 0-byte file was accepted"
    assert result["problem"] == "empty", f"Wrong problem code: {result['problem']}"


def test_missing_file_is_refused():
    result = extract(os.path.join(WORKSPACE, "no_such_document.pdf"))
    assert result["ok"] is False, "A non-existent file was accepted"
    assert result["problem"] == "missing", f"Wrong problem code: {result['problem']}"


def test_display_helper_never_raises():
    """The dashboard uses this, and a bad document must not take the page down."""
    assert load_text_or_empty(os.path.join(WORKSPACE, "nope.pdf")) == ""
    scanned = make_pdf("scanned3.pdf", ["Wording on the page. " * 30], scanned=True)
    assert load_text_or_empty(scanned) == "", "A scanned PDF should display as empty, not raise"


# ----------------------------------------------------------------------------
# Text tidying
# ----------------------------------------------------------------------------

def test_normalize_keeps_the_words_and_hyphens():
    """
    Tidying must not silently rewrite the contract.

    In particular a word hyphenated across a line break is LEFT ALONE, because
    joining "thirty-\\nsix" into "thirtysix" would change what the document says.
    """
    messy = "Section 1.1\r\n\r\n\r\nThe term is thirty-\nsix months from the\tEffective Date.\f"
    tidy = normalize_text(messy)
    assert "thirty-\nsix" in tidy, "A hyphenated line break was rewritten"
    assert "\r" not in tidy and "\f" not in tidy, "Carriage returns / form feeds survived"
    assert " " not in tidy, "A non-breaking space survived"
    assert "\n\n\n" not in tidy, "Runs of blank lines were not collapsed"
    assert "the Effective Date" in tidy, "A tab was not collapsed to a space"


def test_describe_reads_like_a_sentence():
    path = make_pdf("described.pdf", [f"Section 14.1: {CLAUSE}. " + ("Terms follow. " * 30)])
    note = describe(extract(path))
    assert "1-page PDF" in note, f"The provenance note is wrong: {note}"
    assert "pypdf" in note, f"The provenance note should name the engine: {note}"
    assert "characters" in note, f"The provenance note should give a size: {note}"


if __name__ == "__main__":
    tests = [
        test_single_page_pdf_is_read,
        test_multi_page_pdf_reads_every_page,
        test_plain_text_still_works,
        test_non_utf8_text_file_is_still_read,
        test_docx_including_table_text,
        test_scanned_pdf_is_detected_and_not_reviewed,
        test_scanned_pdf_raises_rather_than_returning_empty_text,
        test_partly_scanned_pdf_is_reviewed_but_warns,
        test_tiny_amount_of_text_is_treated_as_unreadable,
        test_damaged_pdf_does_not_crash,
        test_password_protected_pdf_says_so,
        test_unsupported_format_is_refused_by_name,
        test_empty_file_is_refused,
        test_missing_file_is_refused,
        test_display_helper_never_raises,
        test_normalize_keeps_the_words_and_hyphens,
        test_describe_reads_like_a_sentence,
    ]

    WORKSPACE = tempfile.mkdtemp(prefix="covenant_intake_")
    failures = 0
    try:
        for test in tests:
            try:
                test()
                print(f"PASS: {test.__name__}")
            except AssertionError as e:
                failures += 1
                print(f"FAIL: {test.__name__} -- {e}")
            except Exception as e:  # a crash is a failure too, and must be visible
                failures += 1
                print(f"ERROR: {test.__name__} -- {type(e).__name__}: {e}")
    finally:
        shutil.rmtree(WORKSPACE, ignore_errors=True)

    print()
    if failures == 0:
        print("ALL TESTS PASSED")
    else:
        print(f"{failures} TEST(S) FAILED")
        sys.exit(1)
