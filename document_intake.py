"""
DOCUMENT INTAKE — turning a real-world file into text the agents can review.

Why this exists: no law firm keeps its contracts as .txt files. They arrive as
PDFs — and often as PDFs that were scanned from paper, with no text in them at
all. This module is the one place that turns a file on disk into reviewable
text, so every part of the system (the inbox watcher, the orchestrator, the
dashboard's upload feature) reads documents the same way.

What it handles:
  - .pdf   real text extraction, page by page (pypdf, with pdfplumber as a
           fallback if pypdf isn't installed)
  - .docx  Word documents, including text inside tables — contract terms very
           often live in a table
  - .txt   plain text, with a tolerant encoding fallback, because a real .txt
           off a Windows machine is frequently not UTF-8

THE IMPORTANT PART — knowing when we CAN'T read a document:

A scanned PDF is just pictures of pages. Text extraction returns nothing, or a
few characters of noise. Handing that to an AI reviewer would produce a
confident review of an empty document, which is far worse than an error. So
extraction reports a clear, honest failure instead:

  - no text at all on any page      -> "scanned/image-based, OCR isn't built yet"
  - a trace of text but not enough  -> same conclusion, with the character count
  - SOME pages empty, others fine   -> the review goes ahead, but the result
                                       carries a warning naming the blank pages,
                                       so a reviewer knows part of the document
                                       was never read
  - password-protected              -> said plainly, not reported as "empty"
  - damaged / not really a PDF      -> said plainly, and never raised as a crash

Nothing here calls an AI model, and nothing here writes to the file it reads.
Extraction is read-only.

Try it on any file from the command line:
    python document_intake.py samples/jv_texas_refinery.pdf
"""

import os
import re

# Extensions we can turn into text. Anything else is refused by name rather
# than being read as bytes and reviewed as garbage.
SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")

# A generous ceiling for an uploaded contract. Present so a mis-drop (a video,
# a disk image) is refused politely instead of being loaded into memory.
MAX_UPLOAD_BYTES = 25 * 1024 * 1024  # 25 MB

# How we decide a page carried real text rather than a stray mark. A scanner
# stamp or a page number can leave a handful of characters behind, so "not
# empty" is not the same as "has content".
MIN_CHARS_PER_PAGE = 25

# Below this, across the whole document, there is not enough text to review
# honestly — a one-line contract does not exist. Our shortest real sample
# clause is ~200 characters, so this leaves comfortable room.
MIN_TOTAL_CHARS = 120

# The message the user asked for, kept in one place so the dashboard, the
# watcher and the command line all say exactly the same thing.
SCANNED_MESSAGE = (
    "This PDF appears to be scanned/image-based; no text could be extracted. "
    "OCR support isn't built yet."
)


class UnreadableDocument(Exception):
    """
    Raised by load_text() when a document cannot be turned into reviewable text.

    Carries the human-readable explanation as its message, plus a short
    `problem` code for callers that want to react differently to, say, a
    scanned PDF versus a password-protected one.
    """

    def __init__(self, message: str, problem: str = "unreadable"):
        super().__init__(message)
        self.problem = problem


def _result(ok: bool, text: str = "", **fields) -> dict:
    """
    Builds the extraction result in one consistent shape, so callers never have
    to check whether a key exists.

    A plain dict rather than a class, to match how the agents in this project
    already pass their results around.
    """
    result = {
        "ok": ok,
        "text": text,
        "chars": len(text.strip()),
        "format": "",
        "method": "",
        "page_count": 0,
        "pages_with_text": 0,
        "problem": None,
        "message": "",
        "warnings": [],
    }
    result.update(fields)
    return result


# ----------------------------------------------------------------------------
# Cleaning up extracted text
# ----------------------------------------------------------------------------

def normalize_text(text: str) -> str:
    """
    Tidies extracted text without changing a single word of it.

    PDF extraction produces hard line breaks mid-sentence, form feeds between
    pages, and typographic characters that a Windows console can't print. Those
    are all cleaned up. What is deliberately NOT done: joining words that were
    hyphenated across a line break. "thirty-\\nsix" could legitimately be
    "thirty-six" or a split "thirtysix", and guessing wrong in a contract
    changes its meaning. Faithfulness beats tidiness here.
    """
    if not text:
        return ""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\f", "\n")           # page-break characters
    text = text.replace(" ", " ")        # non-breaking space
    text = text.replace("‑", "-")        # non-breaking hyphen
    text = text.replace("ﬁ", "fi").replace("ﬂ", "fl")  # ligatures
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)    # collapse runs of blank lines
    return text.strip()


def _meaningful(page_text: str) -> bool:
    """True if a page carried real text rather than a stray mark or page number."""
    return len(re.sub(r"\s+", "", page_text or "")) >= MIN_CHARS_PER_PAGE


# ----------------------------------------------------------------------------
# PDF
# ----------------------------------------------------------------------------

def _pages_with_pypdf(path: str) -> list:
    """Text of each page, using pypdf. Raises on an unreadable file."""
    from pypdf import PdfReader

    reader = PdfReader(path)
    if reader.is_encrypted:
        # Plenty of real PDFs are "encrypted" with an empty password purely to
        # set permissions, and those open fine. A genuine password stops us.
        try:
            if not reader.decrypt(""):
                raise UnreadableDocument(
                    "This PDF is password-protected, so its text can't be read. "
                    "Please supply an unlocked copy.",
                    "encrypted",
                )
        except UnreadableDocument:
            raise
        except Exception:
            raise UnreadableDocument(
                "This PDF is password-protected, so its text can't be read. "
                "Please supply an unlocked copy.",
                "encrypted",
            )
    return [(page.extract_text() or "") for page in reader.pages]


def _pages_with_pdfplumber(path: str) -> list:
    """Text of each page, using pdfplumber. The fallback engine."""
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        return [(page.extract_text() or "") for page in pdf.pages]


# pypdf is the primary engine, chosen by measuring both on this project's own
# sample contracts (samples/, built by make_test_documents.py). The two returned
# byte-identical text on every one of them, so there was nothing to choose on
# accuracy — but pypdf was an order of magnitude faster (15ms vs 156ms on the
# one-page contract, 106ms vs 2,015ms on the multi-page one), which matters for
# a dashboard where someone is waiting on an upload. pdfplumber stays wired up
# as a fallback so a machine that has only one of the two still works; it is
# also the better tool if this ever needs table-aware layout extraction.
PDF_ENGINES = (("pypdf", _pages_with_pypdf), ("pdfplumber", _pages_with_pdfplumber))


def _extract_pdf(path: str) -> dict:
    """
    Reads a PDF page by page and reports honestly on what came back.

    Multi-page documents are handled by design: every page is extracted
    separately and we keep count of how many of them actually contained text.
    That per-page count is what lets us tell a normal contract apart from a
    scanned one, and from the awkward middle case where only some pages were
    scanned.
    """
    pages = None
    method = ""
    missing_libraries = []

    for name, engine in PDF_ENGINES:
        try:
            pages = engine(path)
            method = name
            break
        except ImportError:
            missing_libraries.append(name)
            continue
        except UnreadableDocument as e:
            return _result(False, format="PDF", method=name, problem=e.problem, message=str(e))
        except Exception as e:
            # A malformed or truncated PDF must never take the app down with it.
            return _result(
                False,
                format="PDF",
                method=name,
                problem="damaged",
                message=(
                    "This file could not be opened as a PDF — it may be damaged "
                    f"or not really a PDF. ({type(e).__name__})"
                ),
            )

    if pages is None:
        return _result(
            False,
            format="PDF",
            problem="library_missing",
            message=(
                "PDF support needs a text-extraction library. Install one with: "
                "pip install pypdf   (tried: " + ", ".join(missing_libraries) + ")"
            ),
        )

    page_count = len(pages)
    if page_count == 0:
        return _result(
            False,
            format="PDF",
            method=method,
            problem="damaged",
            message="This PDF contains no pages at all.",
        )

    readable = [index for index, page in enumerate(pages, start=1) if _meaningful(page)]
    text = normalize_text("\n\n".join(pages))
    chars = len(text.strip())

    base = {
        "format": "PDF",
        "method": method,
        "page_count": page_count,
        "pages_with_text": len(readable),
    }

    # The case a real firm hits constantly: an old contract that was scanned.
    if not readable or chars < MIN_TOTAL_CHARS:
        detail = ""
        if chars:
            detail = (
                f" Only {chars} characters came back from {page_count} "
                f"page{'s' if page_count != 1 else ''}, which is too little to "
                "review as a contract."
            )
        return _result(False, problem="scanned", message=SCANNED_MESSAGE + detail, **base)

    warnings = []
    if len(readable) < page_count:
        # Partly scanned: worth reviewing, but the reviewer must be told that
        # part of the document was never actually read.
        blank = [str(n) for n in range(1, page_count + 1) if n not in readable]
        shown = ", ".join(blank[:8]) + (" …" if len(blank) > 8 else "")
        warnings.append(
            f"{len(blank)} of {page_count} pages produced no text (page{'s' if len(blank) != 1 else ''} "
            f"{shown}) and may be scanned images. Only the text that could be "
            "extracted was reviewed."
        )

    return _result(True, text, warnings=warnings, **base)


# ----------------------------------------------------------------------------
# Word documents
# ----------------------------------------------------------------------------

def _extract_docx(path: str) -> dict:
    """
    Reads a .docx, including text inside tables.

    Tables matter: contract schedules, definitions and payment terms are very
    often laid out as tables, and a reader that only walked the paragraphs
    would silently skip exactly the clauses that get argued about.
    """
    try:
        import docx  # python-docx
    except ImportError:
        return _result(
            False,
            format="Word document",
            problem="library_missing",
            message=(
                "Word support needs the python-docx library. Install it with: "
                "pip install python-docx"
            ),
        )

    try:
        document = docx.Document(path)
    except Exception as e:
        return _result(
            False,
            format="Word document",
            problem="damaged",
            message=(
                "This file could not be opened as a Word document — it may be "
                f"damaged, or an older .doc file rather than .docx. ({type(e).__name__})"
            ),
        )

    blocks = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    text = normalize_text("\n".join(blocks))
    base = {"format": "Word document", "method": "python-docx"}

    if len(text.strip()) < MIN_TOTAL_CHARS:
        return _result(
            False,
            problem="empty",
            message=(
                "This Word document contains almost no text "
                f"({len(text.strip())} characters). If the contract is an image "
                "pasted into the document, OCR support isn't built yet."
            ),
            **base,
        )
    return _result(True, text, **base)


# ----------------------------------------------------------------------------
# Plain text
# ----------------------------------------------------------------------------

def _extract_txt(path: str) -> dict:
    """
    Reads a plain-text document.

    Tries UTF-8 first, then the encodings a real Windows-authored .txt actually
    turns up as. The last attempt cannot fail, so a stray byte never stops a
    document being reviewed.
    """
    raw = None
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except OSError as e:
        return _result(
            False,
            format="plain text",
            problem="missing",
            message=f"This file could not be opened. ({e.strerror or type(e).__name__})",
        )

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any byte sequence
        text = raw.decode("utf-8", errors="replace")
        encoding = "utf-8 (with unreadable characters replaced)"

    text = normalize_text(text)
    base = {"format": "plain text", "method": encoding}
    if not text.strip():
        return _result(False, problem="empty", message="This text file is empty.", **base)
    return _result(True, text, **base)


# ----------------------------------------------------------------------------
# The entry points
# ----------------------------------------------------------------------------

def extract(path: str) -> dict:
    """
    Turns a file into text and describes what happened.

    Always returns a result dict — it does not raise for an unreadable
    document, because "we couldn't read this, and here's why" is a normal
    outcome that the caller needs to show to a human, not an error condition.
    """
    if not os.path.isfile(path):
        return _result(
            False,
            problem="missing",
            message="That file no longer exists on disk.",
        )

    extension = os.path.splitext(path)[1].lower()
    if extension not in SUPPORTED_EXTENSIONS:
        return _result(
            False,
            problem="unsupported",
            message=(
                f"'{extension or 'no extension'}' files aren't supported yet. "
                "Supported formats: " + ", ".join(SUPPORTED_EXTENSIONS) + "."
            ),
        )

    size = os.path.getsize(path)
    if size == 0:
        return _result(False, problem="empty", message="That file is empty (0 bytes).")
    if size > MAX_UPLOAD_BYTES:
        return _result(
            False,
            problem="too_large",
            message=(
                f"That file is {size / 1_048_576:.1f} MB, over the "
                f"{MAX_UPLOAD_BYTES // 1_048_576} MB limit for a single review."
            ),
        )

    if extension == ".pdf":
        return _extract_pdf(path)
    if extension == ".docx":
        return _extract_docx(path)
    return _extract_txt(path)


def load_text(path: str) -> str:
    """
    The simple form, for code that just wants the words: returns the text, or
    raises UnreadableDocument with an explanation a human can act on.

    This is what the Reader agent and the Orchestrator use, which is why a PDF
    dropped into the watched inbox is reviewed properly rather than read as
    binary noise.
    """
    result = extract(path)
    if not result["ok"]:
        raise UnreadableDocument(result["message"], result["problem"] or "unreadable")
    return result["text"]


def load_text_or_empty(path: str) -> str:
    """
    The forgiving form, for display code: returns "" instead of raising.

    The dashboard uses this — a document it can't read should leave a panel
    saying so, not take the page down.
    """
    try:
        return load_text(path)
    except UnreadableDocument:
        return ""


def describe(result: dict) -> str:
    """
    One short sentence recording where a document's text came from, e.g.
    "3-page PDF, text extracted with pypdf (4,812 characters)".

    This gets stored with the review and shown on the case, because a
    compliance team reviewing a redline needs to know whether the text under
    discussion was read cleanly or scraped off a bad scan.
    """
    if not result.get("ok"):
        return result.get("message", "This document could not be read.")

    fmt = result.get("format") or "document"
    pages = result.get("page_count") or 0
    lead = f"{pages}-page {fmt}" if pages else fmt
    verb = "read as" if fmt == "plain text" else "text extracted with"
    note = f"{lead}, {verb} {result.get('method') or 'an unknown method'}"
    note += f" ({result.get('chars', 0):,} characters)"
    if result.get("pages_with_text") and pages and result["pages_with_text"] < pages:
        note += f"; {result['pages_with_text']} of {pages} pages contained text"
    return note


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        print(__doc__)
        print("Usage: python document_intake.py <file>")
        raise SystemExit(0)

    target = sys.argv[1]
    outcome = extract(target)

    print(f"=== DOCUMENT INTAKE: {os.path.basename(target)} ===")
    print(f"Readable: {outcome['ok']}")
    print(f"Summary : {describe(outcome)}")
    if not outcome["ok"]:
        print(f"Problem : {outcome['problem']}")
        print(f"Message : {outcome['message']}")
    for warning in outcome["warnings"]:
        print(f"Warning : {warning}")
    if outcome["ok"]:
        preview = outcome["text"][:600]
        print("--- first 600 characters ---")
        print(preview + ("..." if len(outcome["text"]) > 600 else ""))
